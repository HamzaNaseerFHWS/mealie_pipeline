"""Reads a CSV or XLSX of recipe names into rows (pure apart from opening the file).

Two layouts:

* **list**: a header row (`name`, also `recipe`, `title`, `rezept`; optional `url`, `language`, `notes`,
  `cuisine`) and one recipe per row, or a plain single column of names with no header.
* **grid**: the first row holds category titles, one per column, and *every* cell below it, in every
  column, is one standalone recipe name. Columns may have different lengths (empty cells are skipped).

`auto` picks grid when the first row has no recognised header and the sheet has several columns.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

Layout = Literal["list", "grid"]
LAYOUT_CHOICES = ("auto", "list", "grid")

NAME_HEADERS = ("name", "recipe", "title", "rezept")
OTHER_HEADERS = ("url", "language", "notes", "cuisine")
ACCEPTED_HEADERS = ", ".join(NAME_HEADERS)
# German spellings (a German Excel) map onto the same optional columns.
HEADER_ALIASES = {
    "link": "url",
    "sprache": "language",
    "notizen": "notes",
    "notiz": "notes",
    "küche": "cuisine",
    "kueche": "cuisine",
}
LANGUAGES = {
    "en": "en",
    "english": "en",
    "de": "de",
    "german": "de",
    "deutsch": "de",
    "ur": "ur",
    "urdu": "ur",
    "hi": "hi",
    "hindi": "hi",
}
DELIMITERS = (",", ";", "\t")


class SheetError(Exception):
    """The file cannot be used; the message says why and what to do."""


@dataclass(frozen=True)
class SheetRow:
    row_number: int  # the row as the user sees it in Excel (the header is row 1)
    name: str
    url: str | None = None
    language: str | None = None
    notes: str | None = None
    cuisine: str | None = None
    category: str | None = None  # grid layout: the title of the column the name sits in
    column: int = 1  # 1-based
    position: str = ""  # "R2C1", or "Sheet!R2C1" for a worksheet: points back to the cell
    layout: Layout = "list"


@dataclass
class ParsedSheet:
    layout: Layout = "list"
    rows: list[SheetRow] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)  # list layout: names repeated within the sheet
    repeated: list[str] = field(default_factory=list)  # grid layout: names that sit in several cells
    warnings: list[str] = field(default_factory=list)

    @property
    def columns(self) -> int:
        return len({r.column for r in self.rows})


def normalize_name(name: str) -> str:
    """Case-, space- and punctuation-insensitive identity of a recipe name."""
    return " ".join(re.sub(r"[^\w\s]", " ", name.casefold()).split())


def normalize_url(url: str) -> str:
    cleaned = url.strip().split("#")[0].rstrip("/")
    return re.sub(r"^https?://(www\.)?", "", cleaned, flags=re.IGNORECASE).casefold()


def cell_position(row: int, column: int, sheet: str | None = None) -> str:
    base = f"R{row}C{column}"
    return f"{sheet}!{base}" if sheet else base


def _canonical_headers(first_row: list[str]) -> list[str]:
    return [HEADER_ALIASES.get(c.casefold(), c.casefold()) for c in first_row]


def detect_layout(grid: list[list[str]]) -> Layout:
    """`grid` when the first row is category titles over several columns; otherwise `list`."""
    cleaned = [[c.strip() for c in row] for row in grid]
    first = next((row for row in cleaned if any(row)), None)
    if first is None:
        return "list"
    headers = _canonical_headers(first)
    if any(h in NAME_HEADERS or h in OTHER_HEADERS for h in headers):
        return "list"  # a header row the list layout understands (even if it lacks a name column)
    filled_titles = sum(1 for c in first if c)
    return "grid" if filled_titles >= 2 else "list"


def parse_cells(grid: list[list[str]], layout: str = "auto", sheet: str | None = None) -> ParsedSheet:
    """Turns rows of cell text into `SheetRow`s; row numbers count from 1 over the original grid."""
    if layout not in LAYOUT_CHOICES:
        raise SheetError(f"unknown layout {layout!r}: use one of {', '.join(LAYOUT_CHOICES)}")
    numbered = [(i, [c.strip() for c in row]) for i, row in enumerate(grid, start=1)]
    if not any(any(row) for _, row in numbered):
        raise SheetError("the sheet is empty")
    chosen = detect_layout(grid) if layout == "auto" else layout
    if chosen == "grid":
        return _parse_grid(numbered, sheet)
    return _parse_list(numbered, sheet)


def _parse_grid(numbered: list[tuple[int, list[str]]], sheet: str | None) -> ParsedSheet:
    header_index, header = next((i, row) for i, row in numbered if any(row))
    result = ParsedSheet(layout="grid")
    counts: dict[str, int] = {}
    first_spelling: dict[str, str] = {}
    for index, row in numbered:
        if index <= header_index:
            continue
        for col, name in enumerate(row, start=1):
            if not name or name.startswith("#"):
                continue
            category = header[col - 1] if col - 1 < len(header) and header[col - 1] else None
            result.rows.append(
                SheetRow(
                    index,
                    name,
                    category=category,
                    column=col,
                    position=cell_position(index, col, sheet),
                    layout="grid",
                )
            )
            key = normalize_name(name)
            counts[key] = counts.get(key, 0) + 1
            first_spelling.setdefault(key, name)
    result.repeated = [first_spelling[k] for k, n in counts.items() if n > 1]
    if not result.rows:
        result.warnings.append("the grid has category titles but no recipe names below them")
    return result


def _parse_list(numbered: list[tuple[int, list[str]]], sheet: str | None) -> ParsedSheet:
    first_index, first_row = next((i, row) for i, row in numbered if any(row))
    headers = _canonical_headers(first_row)
    result = ParsedSheet(layout="list")

    columns: dict[str, int]
    if any(h in NAME_HEADERS for h in headers):
        columns = {h: n for n, h in reversed(list(enumerate(headers))) if h}
        name_col = next(columns[h] for h in NAME_HEADERS if h in columns)
        body = [(i, row) for i, row in numbered if i > first_index]
    elif any(h in OTHER_HEADERS for h in headers) or max(len(r) for _, r in numbered) > 1:
        raise SheetError(
            f"no name column found. Put a header row first with one of: {ACCEPTED_HEADERS} "
            "(optional columns: url, language, notes, cuisine), or use --layout grid for a sheet of "
            "category titles with recipe names below them"
        )
    else:  # a plain, headerless list of names
        columns, name_col = {}, 0
        body = [(i, row) for i, row in numbered if i >= first_index]

    def cell(row: list[str], header: str) -> str | None:
        n = columns.get(header)
        return row[n] if n is not None and n < len(row) and row[n] else None

    seen: set[str] = set()
    for index, row in body:
        if not any(row) or (row and row[0].startswith("#")):
            continue
        name = row[name_col] if name_col < len(row) else ""
        if not name:
            result.warnings.append(f"row {index}: no name, skipped")
            continue
        if name.startswith("#"):
            continue
        key = normalize_name(name)
        if key in seen:
            result.duplicates.append(name)
            continue
        seen.add(key)
        url, language = cell(row, "url"), cell(row, "language")
        if url and not re.match(r"https?://", url, re.IGNORECASE):
            result.warnings.append(f"row {index}: {url!r} is not an http(s) URL, ignored")
            url = None
        if language:
            normalized = LANGUAGES.get(language.casefold())
            if normalized is None:
                result.warnings.append(
                    f"row {index}: unknown language {language!r}, ignored (use en, de, ur, hi)"
                )
            language = normalized
        result.rows.append(
            SheetRow(
                index,
                name,
                url,
                language,
                cell(row, "notes"),
                cell(row, "cuisine"),
                column=name_col + 1,
                position=cell_position(index, name_col + 1, sheet),
            )
        )
    return result


def _split_csv(text: str) -> list[list[str]]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    sample = lines[:5] or [""]

    def width(delimiter: str) -> int:
        return max((len(next(csv.reader([ln], delimiter=delimiter))) for ln in sample), default=1)

    delimiter = max(DELIMITERS, key=lambda d: (width(d), d == ","))  # ties go to the comma
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_grid(path: Path, sheet: str | None = None) -> tuple[list[list[str]], str | None]:
    """The cell text of a file, and the worksheet name for an Excel file (None for CSV)."""
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise SheetError(
            f"{path.name} is an old .xls file, which is not supported: save it as .xlsx or .csv and try again"
        )
    if suffix == ".xlsx":
        return _read_xlsx(path, sheet)
    if suffix in (".csv", ".tsv", ".txt"):
        try:
            text = path.read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError as err:
            raise SheetError(f"{path.name} is not UTF-8: save it as 'CSV UTF-8' from Excel") from err
        return _split_csv(text), None
    raise SheetError(f"{path.name}: unsupported file type {suffix!r}. Use .xlsx or .csv")


def _read_xlsx(path: Path, sheet: str | None) -> tuple[list[list[str]], str | None]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as err:  # openpyxl raises many types for a corrupt or non-xlsx file
        raise SheetError(f"cannot read {path.name} as an Excel file: {err}") from err
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            names = ", ".join(workbook.sheetnames)
            raise SheetError(f"no sheet named {sheet!r}. Sheets in the file: {names}")
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        # min_row/min_col keep positions true to the sheet even if the data does not start at A1.
        rows = worksheet.iter_rows(min_row=1, min_col=1, values_only=True)
        grid = [["" if v is None else str(v) for v in row] for row in rows]
        return grid, worksheet.title
    finally:
        workbook.close()


def read_sheet(path: Path, sheet: str | None = None, layout: str = "auto") -> ParsedSheet:
    grid, title = read_grid(path, sheet)
    return parse_cells(grid, layout, title)
