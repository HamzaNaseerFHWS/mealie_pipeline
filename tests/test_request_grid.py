"""The grid layout: category titles in row 1, a recipe name in every cell below, in every column."""

import csv
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from mealie_pipeline.cli import app
from mealie_pipeline.recipe_requests import RecipeRequests
from mealie_pipeline.request_sheet import SheetError, detect_layout, parse_cells, read_sheet
from mealie_pipeline.store import StateStore
from tests.helpers import fake_adapters

FIXTURE = Path(__file__).parent / "fixtures" / "recipe_grid.csv"
CATEGORIES = [
    "Gerichte – Frühstück / Brunch",
    "Gerichte – Currys / Desi",
    "Getränke",
    "Prep / Saucen / Dips / Condiments",
    "Proteins",
]
FATTEH = "Fatteh Bowl (Chicken, Chickpeas, Hummus)"


# -- layout detection (table) ---------------------------------------------------------------


@pytest.mark.parametrize(
    ("grid", "expected"),
    [
        ([["Getränke", "Proteins"], ["Tee", "Tofu"]], "grid"),  # titles over several columns
        ([["Getränke", "Proteins", ""], ["Tee", "", "x"]], "grid"),  # a trailing empty title cell
        (
            [["", "", ""], ["Getränke", "Proteins", "Extras"], ["Tee", "Tofu", "Oliven"]],
            "grid",
        ),  # blank rows first
        ([["name"], ["Dal"]], "list"),
        ([["Name", "Notes"], ["Dal", "x"]], "list"),
        ([["Rezept", "Sprache"], ["Dal", "de"]], "list"),
        ([["title", "Getränke"], ["Dal", "Tee"]], "list"),  # one recognised header makes it a list
        ([["url", "notes"], ["https://x.com", "hi"]], "list"),  # a list header that lacks a name column
        ([["Dal"], ["Roti"], ["Naan"]], "list"),  # one headerless column
        ([["Getränke"], ["Tee"], ["Wasser"]], "list"),  # one column cannot be told from a name list
        ([["Gerichte", "", ""], ["a", "b", "c"]], "list"),  # only one title: not a grid
        ([], "list"),
    ],
)
def test_layout_auto_detection(grid, expected):
    assert detect_layout(grid) == expected


def test_a_one_column_sheet_can_be_forced_to_grid():
    parsed = parse_cells([["Getränke"], ["Tee"], ["Wasser"]], "grid")

    assert [(r.name, r.category, r.position) for r in parsed.rows] == [
        ("Tee", "Getränke", "R2C1"),
        ("Wasser", "Getränke", "R3C1"),
    ]


def test_a_list_sheet_can_be_forced_to_grid_and_a_grid_to_list():
    assert parse_cells([["name"], ["Dal"]], "grid").rows[0].category == "name"  # "name" is a title now
    with pytest.raises(SheetError, match="--layout grid"):
        parse_cells([["Getränke", "Proteins"], ["Tee", "Tofu"]], "list")


def test_an_unknown_layout_is_an_error():
    with pytest.raises(SheetError, match="unknown layout"):
        parse_cells([["a"]], "wide")


def test_a_title_row_with_no_names_warns():
    parsed = parse_cells([["Getränke", "Proteins"], ["", ""]])

    assert parsed.rows == [] and "no recipe names" in parsed.warnings[0]


# -- the grid fixture, parsed ---------------------------------------------------------------


def test_every_cell_below_the_header_is_one_request_with_its_category_and_position():
    parsed = read_sheet(FIXTURE)

    assert parsed.layout == "grid" and parsed.columns == 5 and len(parsed.rows) == 24
    assert [(r.position, r.name, r.category) for r in parsed.rows[:3]] == [
        ("R2C1", "French Toast", CATEGORIES[0]),
        ("R2C2", "Butter Chicken", CATEGORIES[1]),
        ("R2C3", "Tee", CATEGORIES[2]),
    ]
    per_column = {c: sum(1 for r in parsed.rows if r.column == c) for c in range(1, 6)}
    assert per_column == {1: 6, 2: 6, 3: 3, 4: 4, 5: 5}  # columns have different lengths


def test_the_header_row_is_never_a_recipe_and_empty_cells_are_skipped():
    parsed = read_sheet(FIXTURE)

    names = [r.name for r in parsed.rows]
    assert not set(CATEGORIES) & set(names)
    assert "" not in names and all(r.row_number >= 2 for r in parsed.rows)
    assert "R8C1" not in {r.position for r in parsed.rows}  # an empty cell
    assert "R8C5" in {r.position for r in parsed.rows}


def test_a_quoted_cell_with_commas_stays_one_name_and_a_leading_space_is_trimmed():
    names = {r.position: r.name for r in read_sheet(FIXTURE).rows}

    assert names["R5C2"] == FATTEH
    assert names["R6C1"] == "Schwarzbrot mit Tomaten"


def test_the_same_name_in_several_cells_is_kept_per_cell_with_its_own_category():
    parsed = read_sheet(FIXTURE)

    def cells(name):
        return [(r.position, r.category) for r in parsed.rows if r.name == name]

    assert cells("Baked Oats") == [("R4C1", CATEGORIES[0]), ("R7C3", CATEGORIES[2])]
    assert cells("Haleem") == [("R3C2", CATEGORIES[1]), ("R8C5", CATEGORIES[4])]
    assert [p for p, _ in cells("Butter Chicken")] == ["R2C2", "R6C2"]
    assert [p for p, _ in cells("Grießbrei")] == ["R5C1", "R7C1"]
    assert sorted(parsed.repeated) == ["Baked Oats", "Butter Chicken", "Grießbrei", "Haleem"]
    assert parsed.duplicates == []  # the list-layout duplicate rule does not apply to a grid


def test_a_grid_with_an_untitled_column_still_reads_its_names():
    parsed = parse_cells([["Getränke", "", "Proteins"], ["Tee", "Salz", "Tofu"]])

    assert [(r.name, r.category) for r in parsed.rows] == [
        ("Tee", "Getränke"),
        ("Salz", None),
        ("Tofu", "Proteins"),
    ]


# -- the pipeline seam ----------------------------------------------------------------------


@pytest.fixture(autouse=True)
def clean_environment(monkeypatch):
    for name in ("MEALIE_URL", "MEALIE_TOKEN", "OPENAI_API_KEY", "OPENAI_MAX_USD"):
        monkeypatch.delenv(name, raising=False)


@pytest.fixture
def workspace(tmp_path) -> Path:
    (tmp_path / ".env").write_text("MEALIE_URL=http://fake:9000\nMEALIE_TOKEN=t\n", encoding="utf-8")
    shutil.copy(FIXTURE, tmp_path / "food_list_kategorien.csv.csv")  # the real file is named .csv.csv too
    return tmp_path


def mp(workspace: Path, *args: str):
    return CliRunner().invoke(app, ["--workspace", str(workspace), *args], obj=fake_adapters())


def stored(workspace: Path):
    return RecipeRequests(StateStore(workspace / "pipeline.db")).all()


def test_importing_the_grid_registers_exactly_one_request_per_cell(workspace):
    result = mp(workspace, "requests", "add", "--file", str(workspace / "food_list_kategorien.csv.csv"))

    assert result.exit_code == 0, result.output
    assert "layout: grid (auto-detected); 24 request(s) from 5 column(s)" in result.output
    assert "24 new, 0 already known" in result.output
    assert "4 name(s) appear in more than one cell" in result.output
    requests = stored(workspace)
    assert len(requests) == 24
    assert {r.category for r in requests} == set(CATEGORIES)
    by_cell = {r.position: r for r in requests}
    assert by_cell["R5C2"].name == FATTEH and by_cell["R5C2"].category == CATEGORIES[1]
    assert by_cell["R6C1"].name == "Schwarzbrot mit Tomaten"
    assert {r.status for r in requests} == {"pending"}


def test_reimporting_the_grid_creates_no_duplicates(workspace):
    path = str(workspace / "food_list_kategorien.csv.csv")
    mp(workspace, "requests", "add", "--file", path)

    again = mp(workspace, "requests", "add", "--file", path)

    assert "0 new, 24 already known" in again.output
    assert len(stored(workspace)) == 24


def test_a_cell_changed_in_a_new_version_of_the_sheet_is_added_and_the_rest_is_known(workspace):
    path = workspace / "food_list_kategorien.csv.csv"
    mp(workspace, "requests", "add", "--file", str(path))
    edited = path.read_text(encoding="utf-8").replace("Tee,", "Chai,", 1) + ",,,,Skyr\r\n"
    path.write_text(edited, encoding="utf-8")

    result = mp(workspace, "requests", "add", "--file", str(path))

    assert "2 new, 23 already known" in result.output  # the renamed cell and the new one (25 cells)


def test_the_layout_flag_overrides_the_detection_and_the_output_says_so(workspace):
    path = str(workspace / "food_list_kategorien.csv.csv")

    bad = mp(workspace, "requests", "add", "--file", path, "--layout", "list")
    unknown = mp(workspace, "requests", "add", "--file", path, "--layout", "wide")
    forced = mp(workspace, "requests", "add", "--file", path, "--layout", "grid")

    assert bad.exit_code == 2 and "--layout grid" in bad.output
    assert unknown.exit_code == 2 and "unknown layout" in unknown.output
    assert "layout: grid (chosen with --layout)" in forced.output
    assert len(stored(workspace)) == 24


def test_the_simple_one_column_list_still_works_and_says_list(workspace):
    plain = workspace / "names.csv"
    plain.write_text("Palak Paneer\nChicken Karahi\n", encoding="utf-8")

    result = mp(workspace, "requests", "add", "--file", str(plain))

    assert "layout: list (auto-detected); 2 request(s) from 1 column(s)" in result.output
    assert [(r.name, r.category, r.position) for r in stored(workspace)] == [
        ("Palak Paneer", None, "R1C1"),
        ("Chicken Karahi", None, "R2C1"),
    ]


def test_the_report_and_the_listing_show_category_and_cell(workspace):
    mp(workspace, "requests", "add", "--file", str(workspace / "food_list_kategorien.csv.csv"))

    listing = mp(workspace, "requests", "--show", "30").output

    with (workspace / "reports" / "requests.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 24
    assert list(rows[0])[:5] == ["row_number", "position", "name", "category", "notes"]
    cell = next(r for r in rows if r["position"] == "R5C2")
    assert (cell["name"], cell["category"], cell["status"]) == (FATTEH, CATEGORIES[1], "pending")
    assert "R5C2" in listing and "Gerichte – Currys / Desi" in listing and "Requests: 24" in listing
    short = mp(workspace, "requests", "--show", "5").output
    assert "and 19 more" in short


def test_retry_by_cell_touches_one_cell_and_by_name_every_cell_with_that_name(workspace):
    mp(workspace, "requests", "add", "--file", str(workspace / "food_list_kategorien.csv.csv"))
    db = StateStore(workspace / "pipeline.db").db
    db.execute("UPDATE recipe_requests SET status = 'failed', last_error = 'x' WHERE name = 'Butter Chicken'")
    db.commit()

    one = mp(workspace, "requests", "retry", "R2C2")
    after_one = {r.position: r.status for r in stored(workspace) if r.name == "Butter Chicken"}
    both = mp(workspace, "requests", "retry", "butter chicken")
    after_both = {r.position: r.status for r in stored(workspace) if r.name == "Butter Chicken"}

    assert one.exit_code == 0 and after_one == {"R2C2": "pending", "R6C2": "failed"}
    assert both.exit_code == 0 and after_both == {"R2C2": "pending", "R6C2": "pending"}


# -- XLSX ------------------------------------------------------------------------------------


def grid_rows() -> list[list[str]]:
    return list(csv.reader(FIXTURE.read_text(encoding="utf-8").splitlines()))


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Essen"
    for row in grid_rows():
        first.append([cell or None for cell in row])
    second = workbook.create_sheet("Zweites Blatt")
    second.append(["Dips", "Desserts"])
    second.append(["Hummus", "Kheer"])
    second.append([None, "Ras Malai"])
    workbook.save(path)


def test_xlsx_grid_reads_all_rows_and_columns_of_the_first_sheet(workspace):
    path = workspace / "list.xlsx"
    make_workbook(path)

    parsed = read_sheet(path)

    assert parsed.layout == "grid" and len(parsed.rows) == 24 and parsed.columns == 5
    assert parsed.rows[0].position == "Essen!R2C1"
    assert {r.position: r.name for r in parsed.rows}["Essen!R5C2"] == FATTEH


def test_xlsx_several_sheets_are_imported_one_at_a_time_and_never_merged(workspace):
    path = str(workspace / "list.xlsx")
    make_workbook(workspace / "list.xlsx")

    first = mp(workspace, "requests", "add", "--file", path)
    second = mp(workspace, "requests", "add", "--file", path, "--sheet", "Zweites Blatt")
    again = mp(workspace, "requests", "add", "--file", path, "--sheet", "Zweites Blatt")

    assert "24 new" in first.output
    assert "layout: grid" in second.output and "3 request(s) from 2 column(s)" in second.output
    assert "3 new, 0 already known" in second.output and "0 new, 3 already known" in again.output
    positions = {r.position for r in stored(workspace)}
    assert {"Essen!R2C1", "Zweites Blatt!R2C1", "Zweites Blatt!R3C2"} <= positions
    assert len(stored(workspace)) == 27
    assert mp(workspace, "requests", "add", "--file", path, "--sheet", "Nope").exit_code == 2


def test_xlsx_positions_stay_true_when_the_data_does_not_start_at_a1(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B3"] = "Getränke"
    sheet["C3"] = "Proteins"
    sheet["B4"] = "Tee"
    sheet["C5"] = "Tofu"
    path = tmp_path / "offset.xlsx"
    workbook.save(path)

    parsed = read_sheet(path)

    assert [(r.name, r.category, r.position.split("!")[1]) for r in parsed.rows] == [
        ("Tee", "Getränke", "R4C2"),
        ("Tofu", "Proteins", "R5C3"),
    ]
